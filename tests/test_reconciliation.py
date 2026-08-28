"""Reconciliation ingestion.

The governing constraint: row data never reaches a model. A reconciliation is
thousands of lines of client and third-party invoice detail. Bucketing is
arithmetic and belongs in Python; only the aggregate travels.

The second constraint: an unexplained difference must never be quietly filed
under a flattering category. A line with no remark is UNRECONCILED, not
"timing" — because the reply will otherwise assert a position with nothing
behind it.
"""

import io

import pytest

from backend import reconciliation as recon
from backend.domains import gst


def workbook(rows, headers=None, title=True):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    if title:
        ws.append(["ITC Reconciliation FY 2019-20"])
        ws.append([])
    ws.append(headers or ["Supplier GSTIN", "Supplier Name", "Invoice No",
                          "As per Books", "As per GSTR-2A", "Difference",
                          "Remarks"])
    for row in rows:
        ws.append(list(row))
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


ROWS = [
    ("29AAACX1111A1Z1", "Alpha", "INV-1", 1810000, 0, 1810000,
     "Supplier filed GSTR-1 in subsequent period"),
    ("29AAACX2222B1Z2", "Beta", "INV-2", 150000, 0, 150000,
     "Supplier not filed - return defaulter"),
    ("29AAACX3333C1Z3", "Gamma", "RCM-1", 65000, 0, 65000,
     "RCM - reverse charge self invoice"),
    ("29AAACX4444D1Z4", "Delta", "INV-4", 35000, 0, 35000,
     "Ineligible - blocked credit 17(5)"),
    ("29AAACX5555E1Z5", "Epsilon", "INV-5", 40000, 0, 40000, ""),
]


class TestColumnDetection:
    def test_common_headers(self):
        mapping = recon.detect_columns(
            ["Supplier GSTIN", "Invoice No", "As per Books",
             "As per GSTR-2A", "Difference", "Remarks"]
        )
        for field in ("supplier_gstin", "invoice_no", "amount_books",
                      "amount_2a", "difference", "status"):
            assert field in mapping, f"missed {field}"

    def test_alternative_wording(self):
        mapping = recon.detect_columns(
            ["GSTIN of Supplier", "Party Name", "Bill No", "As per PR",
             "GSTR-2B", "Variance", "Nature of Difference"]
        )
        assert "supplier_gstin" in mapping
        assert "amount_2a" in mapping
        assert "status" in mapping

    def test_case_and_punctuation_insensitive(self):
        assert "supplier_gstin" in recon.detect_columns(["  supplier/gstin  "])

    def test_one_column_used_once(self):
        mapping = recon.detect_columns(["GSTIN", "Amount", "Remarks"])
        assert len(set(mapping.values())) == len(mapping)

    def test_unknown_headers_yield_nothing(self):
        assert recon.detect_columns(["Foo", "Bar", "Baz"]) == {}


class TestParsing:
    def test_finds_header_below_a_title_row(self):
        headers, rows, _ = recon.parse_workbook("r.xlsx", workbook(ROWS))
        assert "Supplier GSTIN" in [str(h) for h in headers]
        assert len(rows) == len(ROWS)

    def test_csv(self):
        content = ("Supplier GSTIN,Difference,Remarks\n"
                   "29AAACX1111A1Z1,1000,Timing\n").encode()
        headers, rows, _ = recon.parse_workbook("r.csv", content)
        assert len(rows) == 1

    def test_semicolon_delimited_csv(self):
        content = ("Supplier GSTIN;Difference;Remarks\n"
                   "29AAACX1111A1Z1;1000;Timing\n").encode()
        _, rows, _ = recon.parse_workbook("r.csv", content)
        assert len(rows) == 1

    def test_rejects_unsupported_type(self):
        with pytest.raises(ValueError, match="Unsupported file type"):
            recon.parse_workbook("r.rtf", b"x" * 100)

    def test_a_pdf_that_is_not_a_pdf_says_so(self):
        """A supported extension over unreadable bytes still raises, and the
        message names the file rather than surfacing a library's internals as
        an unsupported-type error."""
        with pytest.raises(ValueError, match="Could not read the PDF"):
            recon.parse_workbook("r.pdf", b"x" * 100)

    def test_rejects_oversized(self):
        with pytest.raises(ValueError, match="larger than"):
            recon.parse_workbook("r.xlsx", b"x" * (recon.MAX_UPLOAD_BYTES + 1))

    def test_rejects_empty(self):
        with pytest.raises(ValueError, match="no data"):
            recon.parse_workbook("r.csv", b"\n\n")


class TestAmounts:
    def test_indian_formatting(self):
        assert recon._to_amount("41,20,000") == 4120000.0

    def test_currency_symbols_stripped(self):
        assert recon._to_amount("Rs. 1,000.50") == 1000.50

    def test_accounting_negatives(self):
        assert recon._to_amount("(5,000)") == -5000.0

    def test_blank_is_none(self):
        for value in (None, "", "   ", "-"):
            assert recon._to_amount(value) is None

    def test_difference_column_is_authoritative(self):
        mapping = {"difference": 0, "amount_books": 1, "amount_2a": 2}
        assert recon.row_amount([500, 1000, 200], mapping) == 500

    def test_falls_back_to_books_minus_portal(self):
        mapping = {"amount_books": 0, "amount_2a": 1}
        assert recon.row_amount([1000, 200], mapping) == 800

    def test_absolute_value_taken(self):
        mapping = {"amount_books": 0, "amount_2a": 1}
        assert recon.row_amount([200, 1000], mapping) == 800


class TestClassification:
    MAPPING = {"supplier_gstin": 0, "supplier_name": 1, "invoice_no": 2,
               "amount_books": 3, "amount_2a": 4, "difference": 5, "status": 6}

    def _classify(self, remark, name="Supplier", invoice="INV-1"):
        return recon.classify_row(
            ["29AAACX1111A1Z1", name, invoice, 1000, 0, 1000, remark],
            self.MAPPING, gst,
        )

    def test_reads_the_preparers_own_remark(self):
        cases = {
            "Supplier filed GSTR-1 in subsequent period": "timing",
            "Supplier not filed - return defaulter": "non_filer",
            "RCM - reverse charge": "rcm",
            "Ineligible - blocked credit 17(5)": "ineligible",
            "Import IGST - bill of entry": "import_igst",
            "ISD credit distributed": "isd",
            "Credit note amendment": "amendment",
            "Data entry error": "clerical",
            "Supplier wrongly reported under another GSTIN": "supplier_error",
        }
        for remark, expected in cases.items():
            assert self._classify(remark) == expected, f"{remark!r}"

    def test_blank_remark_is_unreconciled_not_flattered(self):
        """The whole point. An unexplained line has no argument behind it."""
        assert self._classify("") == gst.UNRECONCILED.key

    def test_unrecognised_remark_is_unreconciled(self):
        assert self._classify("see annexure 4") == gst.UNRECONCILED.key

    def test_mechanical_category_detected_from_the_reference(self):
        assert self._classify("", invoice="RCM-0012") == "rcm"

    def test_remark_beats_the_reference(self):
        assert self._classify("Timing difference", invoice="RCM-1") == "timing"


class TestSummary:
    def setup_method(self):
        self.summary = recon.read_reconciliation("r.xlsx", workbook(ROWS), gst)

    def test_totals(self):
        assert self.summary["total"] == 2100000.0
        assert self.summary["row_count"] == 5

    def test_buckets_ordered_by_amount(self):
        amounts = [b["amount"] for b in self.summary["buckets"]]
        assert amounts == sorted(amounts, reverse=True)

    def test_largest_bucket_is_timing(self):
        assert self.summary["buckets"][0]["key"] == "timing"

    def test_shares_sum_to_one(self):
        assert sum(b["share"] for b in self.summary["buckets"]) == pytest.approx(1.0)

    def test_each_bucket_carries_its_legal_position(self):
        for bucket in self.summary["buckets"]:
            assert bucket["strength"] in ("strong", "defensible", "weak", "concede")
            assert bucket["label"]
            assert bucket["action"]

    def test_immaterial_rows_excluded(self):
        rows = ROWS + [("29AAACX9999Z1Z9", "Tiny", "INV-9", 0.4, 0, 0.4, "")]
        summary = recon.read_reconciliation("r.xlsx", workbook(rows), gst)
        assert summary["skipped_rows"] == 1

    def test_unreconciled_share_is_warned_about(self):
        rows = [("29AAACX1111A1Z1", "A", "INV-1", 100000, 0, 100000, ""),
                ("29AAACX2222B1Z2", "B", "INV-2", 1000, 0, 1000, "Timing")]
        summary = recon.read_reconciliation("r.xlsx", workbook(rows), gst)
        assert any("unexplained" in w for w in summary["warnings"])

    def test_missing_status_column_is_warned_about(self):
        rows = [("29AAACX1111A1Z1", 1000)]
        content = workbook(rows, headers=["Supplier GSTIN", "Difference"])
        summary = recon.read_reconciliation("r.xlsx", content, gst)
        assert any("remarks column" in w for w in summary["warnings"])

    def test_no_amount_column_is_an_error(self):
        content = workbook([("29AAACX1111A1Z1", "note")],
                           headers=["Supplier GSTIN", "Remarks"])
        with pytest.raises(ValueError, match="No amount column"):
            recon.read_reconciliation("r.xlsx", content, gst)

    def test_unidentifiable_columns_are_an_error(self):
        content = workbook([(1, 2)], headers=["Foo", "Bar"], title=False)
        with pytest.raises(ValueError, match="None of the columns could be identified"):
            recon.read_reconciliation("r.xlsx", content, gst)


class TestSupplierPrivacy:
    """Supplier GSTINs are third-party data, not the client's to disclose."""

    def test_masked_on_the_anonymising_tier(self):
        summary = recon.read_reconciliation(
            "r.xlsx", workbook(ROWS), gst, mask_suppliers=True
        )
        for exposure in summary["top_exposures"]:
            assert exposure["supplier"] == "[supplier withheld]"
            assert "29AAACX" not in exposure["supplier"]

    def test_shown_on_the_pro_tier(self):
        summary = recon.read_reconciliation(
            "r.xlsx", workbook(ROWS), gst, mask_suppliers=False
        )
        assert any("29AAACX" in e["supplier"] for e in summary["top_exposures"])


class TestBriefing:
    """What the panel receives — aggregates only, and small."""

    def setup_method(self):
        self.summary = recon.read_reconciliation("r.xlsx", workbook(ROWS), gst)
        self.brief = gst.reconciliation_brief(self.summary)

    def test_contains_the_buckets_and_amounts(self):
        assert "Timing" in self.brief
        assert "18,10,000" in self.brief or "1,810,000" in self.brief

    def test_carries_the_legal_position_for_each_bucket(self):
        assert "section 16(2)" in self.brief

    def test_instructs_the_panel_to_argue_category_by_category(self):
        assert "single figure" in self.brief or "category by category" in self.brief.lower()

    def test_no_invoice_level_data_travels(self):
        """The constraint that makes this affordable and private."""
        for leak in ("INV-1", "INV-2", "Alpha", "Beta", "29AAACX1111A1Z1"):
            assert leak not in self.brief, f"row data leaked into the brief: {leak}"

    def test_stays_compact(self):
        """Aggregates, not rows — a few hundred tokens however big the file."""
        assert len(self.brief) < 6000

    def test_size_is_independent_of_row_count(self):
        many = [
            (f"29AAACX{i:04d}A1Z1", f"S{i}", f"INV-{i}", 1000, 0, 1000, "Timing")
            for i in range(500)
        ]
        big = gst.reconciliation_brief(
            recon.read_reconciliation("r.xlsx", workbook(many), gst)
        )
        assert len(big) < 6000

    def test_empty_summary_yields_nothing(self):
        assert gst.reconciliation_brief({}) == ""
        assert gst.reconciliation_brief(None) == ""


class TestBucketDefinitions:
    def test_every_bucket_is_complete(self):
        for bucket in gst.RECONCILIATION_BUCKETS:
            assert bucket.key and bucket.label and bucket.position and bucket.action
            assert bucket.strength in ("strong", "defensible", "weak", "concede")

    def test_keys_are_unique(self):
        keys = [b.key for b in gst.RECONCILIATION_BUCKETS]
        assert len(keys) == len(set(keys))

    def test_non_filer_is_marked_weak(self):
        """s.16(2)(c) is the exposed limb and must not read as comfortable."""
        assert gst.RECONCILIATION_BUCKETS_BY_KEY["non_filer"].strength == "weak"

    def test_mechanical_categories_are_strong(self):
        for key in ("rcm", "import_igst", "isd"):
            assert gst.RECONCILIATION_BUCKETS_BY_KEY[key].strength == "strong"

    def test_blocked_credit_is_a_concession(self):
        assert gst.RECONCILIATION_BUCKETS_BY_KEY["ineligible"].strength == "concede"


def _docx_bytes(rows):
    """A Word document holding one table, as a client's ERP export would."""
    import io as _io
    import docx
    document = docx.Document()
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            table.cell(r, c).text = str(value)
    buffer = _io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pdf_bytes(lines):
    """
    A minimal one-page PDF with a real text layer, built by hand.

    Written out rather than pulled from a fixture file so the test says what
    it is testing: a PDF that records where its glyphs are and nothing about
    what its columns were, which is the case the splitter has to survive.
    """
    text = "\n".join(
        f"BT /F1 10 Tf 40 {760 - i * 14} Td ({line}) Tj ET"
        for i, line in enumerate(lines)
    )
    objects = [
        "<</Type/Catalog/Pages 2 0 R>>",
        "<</Type/Pages/Kids[3 0 R]/Count 1>>",
        "<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]"
        "/Resources<</Font<</F1 5 0 R>>>>/Contents 4 0 R>>",
        f"<</Length {len(text)}>>\nstream\n{text}\nendstream",
        "<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (f"trailer<</Size {len(objects) + 1}/Root 1 0 R>>\n"
            f"startxref\n{xref}\n%%EOF").encode()
    return bytes(out)


class TestWordReconciliations:
    """
    A reconciliation reaches the firm in whatever the client's accountant had
    to hand, and that is routinely a Word table pasted out of the ERP.
    Refusing it meant re-keying thousands of rows into a spreadsheet before
    the panel could see any figures — which meant, in practice, it never did.

    A .docx carries its tables as real structure, so this path is exact.
    """

    ROWS = [
        ["Supplier GSTIN", "Supplier Name", "Difference", "Remarks"],
        ["29AAACX1111A1Z1", "Alpha Traders", "12000", "Timing — filed in April"],
        ["29AAACX2222A1Z2", "Beta Supplies", "8000", "Supplier not filed"],
    ]

    def test_a_word_table_is_read(self):
        headers, rows, _ = recon.parse_workbook("r.docx", _docx_bytes(self.ROWS))
        assert len(rows) == 2
        assert "Supplier GSTIN" in headers

    def test_the_figures_survive_the_round_trip(self):
        _, rows, _ = recon.parse_workbook("r.docx", _docx_bytes(self.ROWS))
        assert "12000" in [str(cell) for cell in rows[0]]

    def test_a_document_with_no_table_is_refused(self):
        import io as _io
        import docx
        document = docx.Document()
        document.add_paragraph("The reconciliation is attached separately.")
        buffer = _io.BytesIO()
        document.save(buffer)
        with pytest.raises(ValueError, match="No table was found"):
            recon.parse_workbook("r.docx", buffer.getvalue())

    def test_several_tables_warn_rather_than_merge(self):
        """
        A reconciliation exported to Word usually carries a small summary
        table beside the detail. Stacking the two double-counts every bucket
        it touches, so the widest is read and the rest are named.
        """
        import io as _io
        import docx
        document = docx.Document()
        summary = document.add_table(rows=2, cols=2)
        summary.cell(0, 0).text = "Bucket"
        summary.cell(0, 1).text = "Amount"
        detail = document.add_table(rows=len(self.ROWS), cols=4)
        for r, row in enumerate(self.ROWS):
            for c, value in enumerate(row):
                detail.cell(r, c).text = str(value)
        buffer = _io.BytesIO()
        document.save(buffer)
        _, rows, warnings = recon.parse_workbook("r.docx", buffer.getvalue())
        assert len(rows) == 2
        assert any("2 tables" in w for w in warnings)


class TestPdfReconciliations:
    """
    A PDF records where the glyphs are, not what the columns were, so this
    path checks its own work: the rows must agree on how many columns they
    have, and a file that cannot produce a consistent grid raises and names
    the formats that will work.

    A bucket total computed from a misaligned grid arrives looking exactly
    like a correct one, and is argued from in a filed reply.
    """

    def test_a_scan_with_no_text_layer_is_refused_not_guessed(self):
        """
        OCR is offered for notices because a misread word in a notice is
        visible to the reviewer reading it. A misread digit in row 4,000 of a
        reconciliation is not.
        """
        blank = (b"%PDF-1.4\n1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n"
                 b"2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj\n"
                 b"3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]>>endobj\n"
                 b"trailer<</Root 1 0 R>>")
        with pytest.raises(ValueError, match="no text layer|Could not read"):
            recon.parse_workbook("r.pdf", blank)

    LINES = [
        "Supplier GSTIN    Supplier Name    Difference    Remarks",
        "29AAACX1111A1Z1    Alpha Traders    12000    Timing",
        "29AAACX2222A1Z2    Beta Supplies    8000    Not filed",
    ]

    def test_a_printed_grid_is_recovered(self):
        headers, rows, _ = recon.parse_workbook("r.pdf", _pdf_bytes(self.LINES))
        assert [str(h) for h in headers][:2] == [
            "Supplier GSTIN", "Supplier Name"]
        assert len(rows) == 2
        assert rows[0][2] == "12000"

    def test_page_furniture_is_dropped_and_reported(self):
        """A heading or a page number does not fit the grid. Skipping it
        silently is how a row count stops matching the document."""
        _, rows, warnings = recon.parse_workbook(
            "r.pdf", _pdf_bytes(["ITC Reconciliation FY 2022-23"]
                                + self.LINES + ["Page 1 of 4"]))
        assert len(rows) == 2
        assert any("did not fit" in w for w in warnings)

    def test_a_pdf_with_no_column_separation_is_refused(self):
        with pytest.raises(ValueError, match="No columns|does not appear"):
            recon.parse_workbook(
                "r.pdf", _pdf_bytes(["A letter about the reconciliation.",
                                     "It is attached separately."]))

    def test_a_single_space_does_not_split_a_supplier_name(self):
        """One space is inside a supplier name far more often than it is
        between columns."""
        split = recon._PDF_COLUMN_GAP.split("Alpha Traders Private Limited")
        assert split == ["Alpha Traders Private Limited"]
